import os
import openpyxl
from datetime import datetime, date
from collections import defaultdict
import pymysql
import config  # carga las variables de entorno desde .env

EXCEL_PATH = r'G:\Mi unidad\SISTEMA TECH-HOUSE\TECH HOUSE.xlsm'
TODAY = date.today().isoformat()  # Excluir ventas de hoy

wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

# ==================== 1. LEER ENTRADAS ====================
print("=== Leyendo Entradas ===")
ws_ent = wb['Entrada']
entradas = []
for row in ws_ent.iter_rows(min_row=3, values_only=True):
    if not row[0] or not row[1]:
        continue
    entradas.append({
        'fecha': row[0] if isinstance(row[0], datetime) else datetime(2026, 4, 17),
        'codigo_producto': str(row[1]).strip().upper(),
        'cantidad': int(row[3]) if row[3] else 0,
        'costo_unitario': float(row[4]) if row[4] else 0
    })
print(f"  {len(entradas)} entradas")

# ==================== 2. LEER INVENTARIO (valores cacheados) ====================
print("=== Leyendo Inventario (cacheado) ===")
ws_inv = wb['Inventario']
productos = {}
for row in ws_inv.iter_rows(min_row=3, values_only=True):
    if not row[0] or not str(row[0]).strip().upper().startswith('PROD') and not str(row[0]).strip().startswith(('PROD', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ', 'AK', 'AL', 'AM', 'AN', 'AO', 'AP', 'AQ', 'AR', 'AS', 'AT', 'AU', 'AV', 'AW', 'AX', 'AY', 'AZ', 'CH', 'DH', 'ID', 'MM', 'PT', 'AB')):
        continue
    codigo = str(row[0]).strip().upper()
    nombre = str(row[1]).strip() if row[1] else codigo
    # Cached values from Excel formulas
    costo = float(row[2]) if row[2] is not None else 0
    margen_raw = float(row[3]) if row[3] is not None else 0
    margen = round(margen_raw * 100, 2) if margen_raw < 1 else round(margen_raw, 2)
    precio = float(row[4]) if row[4] is not None else costo
    stock = int(row[5]) if row[5] is not None else 0
    productos[codigo] = {
        'codigo': codigo, 'nombre': nombre, 'costo': costo,
        'margen': margen, 'precio': precio, 'stock': max(0, stock)
    }
print(f"  {len(productos)} productos del Inventario")

# Agregar productos que estan en entradas pero no en inventario
for e in entradas:
    if e['codigo_producto'] not in productos:
        # Buscar nombre en entradas
        nombre = e['codigo_producto']
        for e2 in entradas:
            if e2['codigo_producto'] == e['codigo_producto']:
                # No tenemos nombre aqui, usar codigo
                break
        productos[e['codigo_producto']] = {
            'codigo': e['codigo_producto'], 'nombre': nombre,
            'costo': e['costo_unitario'], 'margen': 0,
            'precio': e['costo_unitario'], 'stock': 0
        }

# Buscar nombres faltantes del Excel (columna C de Entrada = VLOOKUP)
ws_ent2 = wb['Entrada']
for row in ws_ent2.iter_rows(min_row=3, values_only=True):
    if row[1] and row[2]:
        codigo = str(row[1]).strip().upper()
        nombre = str(row[2]).strip()
        if codigo in productos and productos[codigo]['nombre'] == codigo:
            productos[codigo]['nombre'] = nombre

# También del Inventario
ws_inv2 = wb['Inventario']
for row in ws_inv2.iter_rows(min_row=3, values_only=True):
    if row[0] and row[1]:
        codigo = str(row[0]).strip().upper()
        nombre = str(row[1]).strip()
        if codigo in productos:
            productos[codigo]['nombre'] = nombre

print(f"  {len(productos)} productos totales")

# ==================== 3. LEER VENTAS (excluir hoy) ====================
print("=== Leyendo Ventas ===")
ws_ven = wb['Ventas']
ventas_raw = []
excluidas = 0
for row in ws_ven.iter_rows(min_row=3, values_only=True):
    if not row[0] or not row[1]:
        continue
    fecha = row[0] if isinstance(row[0], datetime) else datetime(2026, 1, 1)
    # Excluir ventas de hoy
    if fecha.strftime('%Y-%m-%d') == TODAY:
        excluidas += 1
        continue
    ventas_raw.append({
        'fecha': fecha,
        'codigo': str(row[1]).strip().upper(),
        'producto': str(row[2]).strip(),
        'cantidad': int(row[3]) if row[3] else 1,
        'total': round(float(row[4]) if row[4] else 0, 2)
    })
print(f"  {len(ventas_raw)} ventas ({excluidas} excluidas de hoy)")

# ==================== 4. LEER REPORTE (ganancia directa) ====================
print("=== Leyendo Reporte ===")
ws_rep = wb['Reporte']
ganancia_lookup = defaultdict(list)
for row in ws_rep.iter_rows(min_row=3, values_only=True):
    if row[0] and row[1]:
        key = (str(row[0])[:10], str(row[1]).strip().upper())
        cant = int(row[3]) if row[3] else 0
        ganancia = round(float(row[5]), 2) if row[5] else 0
        ganancia_lookup[key].append((cant, ganancia))
print(f"  {sum(len(v) for v in ganancia_lookup.values())} items con ganancia")

wb.close()

# ==================== 5. AGRUPAR VENTAS ====================
print("\n=== Agrupando Ventas ===")
ventas_agrupadas = defaultdict(list)
for v in ventas_raw:
    date_key = v['fecha'].strftime('%Y-%m-%d')
    ventas_agrupadas[date_key].append(v)

ventas_finales = []
detalle_finales = []
venta_id_counter = 1

for date_key in sorted(ventas_agrupadas.keys()):
    items = ventas_agrupadas[date_key]
    total_venta = round(sum(i['total'] for i in items), 2)
    ganancia_total = 0

    for item in items:
        g_key = (date_key, item['codigo'])
        # Match by date + code + cantidad
        ganancia = 0
        rep_list = ganancia_lookup.get(g_key, [])
        for i, (r_cant, r_gan) in enumerate(rep_list):
            if r_cant == item['cantidad']:
                ganancia = r_gan
                # Remove used entry to avoid double-matching
                ganancia_lookup[g_key] = rep_list[:i] + rep_list[i+1:]
                break
        ganancia_total += ganancia
        precio_unit = round(item['total'] / item['cantidad'], 2) if item['cantidad'] > 0 else item['total']
        detalle_finales.append({
            'venta_id': venta_id_counter,
            'codigo_producto': item['codigo'],
            'cantidad': item['cantidad'],
            'precio_unitario': precio_unit,
            'total': item['total'],
            'ganancia': ganancia
        })

    ventas_finales.append({
        'id': venta_id_counter,
        'fecha': items[0]['fecha'],
        'total_venta': total_venta,
        'ganancia_total': round(ganancia_total, 2)
    })
    venta_id_counter += 1

print(f"  {len(ventas_finales)} ventas")
print(f"  {len(detalle_finales)} detalles")

# Verificar totales
total_v = sum(v['total_venta'] for v in ventas_finales)
total_g = sum(v['ganancia_total'] for v in ventas_finales)
print(f"  Total ventas: Q{total_v:.2f}")
print(f"  Total ganancia: Q{total_g:.2f}")

# ==================== 6. INSERTAR EN MYSQL ====================
print("\n=== Insertando en MySQL ===")
conn = pymysql.connect(host=os.environ.get('MYSQL_HOST', 'localhost'),
                       user=os.environ.get('MYSQL_USER', 'root'),
                       password=os.environ.get('MYSQL_PASSWORD', ''),
                       database=os.environ.get('MYSQL_DB', 'the_tech_house'),
                       charset='utf8mb4')
cursor = conn.cursor()

print("  Limpiando tablas...")
cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
cursor.execute("TRUNCATE TABLE detalle_ventas")
cursor.execute("TRUNCATE TABLE ventas")
cursor.execute("TRUNCATE TABLE entradas")
cursor.execute("TRUNCATE TABLE productos")
cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
conn.commit()

# Productos
print(f"  Insertando {len(productos)} productos...")
sql_prod = "INSERT INTO productos (codigo, nombre, costo, margen, precio, stock) VALUES (%s, %s, %s, %s, %s, %s)"
for p in productos.values():
    cursor.execute(sql_prod, (p['codigo'], p['nombre'], p['costo'], p['margen'], p['precio'], p['stock']))
conn.commit()
cursor.execute("SELECT COUNT(*) FROM productos")
print(f"    OK: {cursor.fetchone()[0]} productos")

# Entradas
print(f"  Insertando {len(entradas)} entradas...")
sql_ent = "INSERT INTO entradas (fecha, codigo_producto, cantidad, costo_unitario) VALUES (%s, %s, %s, %s)"
for e in entradas:
    cursor.execute(sql_ent, (e['fecha'], e['codigo_producto'], e['cantidad'], e['costo_unitario']))
conn.commit()
print(f"    OK")

# Ventas
print(f"  Insertando {len(ventas_finales)} ventas...")
sql_ven = "INSERT INTO ventas (id, fecha, total_venta, ganancia_total) VALUES (%s, %s, %s, %s)"
for v in ventas_finales:
    cursor.execute(sql_ven, (v['id'], v['fecha'], v['total_venta'], v['ganancia_total']))
cursor.execute("ALTER TABLE ventas AUTO_INCREMENT = %s", (venta_id_counter,))
conn.commit()
print(f"    OK")

# Detalles
print(f"  Insertando {len(detalle_finales)} detalles...")
sql_det = "INSERT INTO detalle_ventas (venta_id, codigo_producto, cantidad, precio_unitario, total, ganancia) VALUES (%s, %s, %s, %s, %s, %s)"
for d in detalle_finales:
    cursor.execute(sql_det, (d['venta_id'], d['codigo_producto'], d['cantidad'], d['precio_unitario'], d['total'], d['ganancia']))
conn.commit()
print(f"    OK")

cursor.close()
conn.close()

print("\n=== Migracion completada ===")
print(f"  Productos: {len(productos)}")
print(f"  Entradas: {len(entradas)}")
print(f"  Ventas: {len(ventas_finales)}")
print(f"  Detalles: {len(detalle_finales)}")
print(f"  Total ventas: Q{total_v:.2f}")
print(f"  Total ganancia: Q{total_g:.2f}")
