from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, jsonify
from flask_login import login_required, current_user
from extensions import db
from models import Producto, Entrada, DetalleVenta
from io import BytesIO

registros_bp = Blueprint('registros', __name__)


def _productos_unicos():
    """Devuelve productos sin repetidos, en el orden en que fueron registrados."""
    productos = Producto.query.order_by(Producto.id.asc()).all()
    vistos = set()
    unicos = []
    for p in productos:
        if p.codigo in vistos:
            continue
        vistos.add(p.codigo)
        unicos.append(p)
    return unicos


def _vista():
    return url_for('registros.registros_view')


def _validar_datos(upc, codigo, descripcion, excluir_id=None):
    """Valida los 3 campos. Devuelve mensaje de error o None si todo esta bien."""
    if not codigo or not descripcion:
        return 'El codigo y la descripcion son obligatorios.'
    if len(codigo) > 50:
        return 'El codigo no puede superar 50 caracteres.'
    if len(descripcion) > 150:
        return 'La descripcion no puede superar 150 caracteres.'
    if len(upc) > 50:
        return 'El UPC no puede superar 50 caracteres.'

    q = Producto.query.filter(Producto.codigo == codigo)
    if excluir_id:
        q = q.filter(Producto.id != excluir_id)
    if q.first():
        return f'El codigo "{codigo}" ya existe en otro producto.'

    if upc:
        q = Producto.query.filter(Producto.codigo_barras == upc)
        if excluir_id:
            q = q.filter(Producto.id != excluir_id)
        otro = q.first()
        if otro:
            return f'El UPC "{upc}" ya esta en uso por el producto {otro.codigo}.'
    return None


def identificar_producto(upc, codigo):
    """Busca el producto en la matriz (Registros Nuevos) por UPC y/o codigo.

    Devuelve (producto, mensaje_de_conflicto). Si no existe devuelve (None, None).
    """
    upc = (upc or '').strip()
    codigo = (codigo or '').strip()

    prod_upc = None
    if upc:
        prod_upc = Producto.query.filter(
            (Producto.codigo_barras == upc) | (Producto.codigo == upc)
        ).first()

    prod_cod = None
    if codigo:
        prod_cod = Producto.query.filter_by(codigo=codigo).first()

    if prod_upc and prod_cod and prod_upc.id != prod_cod.id:
        return None, f'El UPC "{upc}" y el codigo "{codigo}" pertenecen a productos distintos.'

    return (prod_upc or prod_cod), None


@registros_bp.route('/api/identificar')
@login_required
def api_identificar():
    """Comprueba en la matriz si el UPC/codigo ya existe (para autocompletar)."""
    if current_user.rol != 'admin':
        return jsonify({'ok': False, 'mensaje': 'Sin acceso.'}), 403

    upc = request.args.get('upc', '').strip()
    codigo = request.args.get('codigo', '').strip()

    if not upc and not codigo:
        return jsonify({'ok': True, 'existe': False, 'parcial': False})

    producto, conflicto = identificar_producto(upc, codigo)
    if conflicto:
        return jsonify({'ok': True, 'existe': False, 'conflicto': True,
                        'mensaje': conflicto})
    if not producto:
        return jsonify({'ok': True, 'existe': False, 'parcial': True})

    return jsonify({
        'ok': True,
        'existe': True,
        'id': producto.id,
        'codigo': producto.codigo,
        'nombre': producto.nombre,
        'upc': producto.codigo_barras or '',
        'margen': round(float(producto.margen), 2),
        'precio': round(float(producto.precio), 2),
    })


@registros_bp.route('/', methods=['GET', 'POST'])
@login_required
def registros_view():
    if current_user.rol != 'admin':
        flash('Solo el administrador puede ver los registros.', 'danger')
        return redirect(url_for('main.index'))

    if request.method == 'POST':
        action = request.form.get('action', 'crear')
        upc = request.form.get('upc', '').strip()
        codigo = request.form.get('codigo', '').strip()
        descripcion = request.form.get('descripcion', '').strip()
        id_prod = request.form.get('id_producto', '').strip()

        # ---- EDITAR ----
        if action == 'editar':
            producto = Producto.query.get(id_prod) if id_prod else None
            if not producto:
                flash('El producto a editar no existe.', 'danger')
                return redirect(_vista())

            error = _validar_datos(upc, codigo, descripcion, excluir_id=producto.id)
            if error:
                flash(error, 'danger')
                return redirect(_vista())

            codigo_anterior = producto.codigo
            producto.codigo = codigo
            producto.nombre = descripcion
            producto.codigo_barras = upc if upc else None

            # El FK de entradas propaga solo (ON UPDATE CASCADE),
            # pero detalle_ventas no tiene FK: se actualiza a mano.
            if codigo_anterior != codigo:
                DetalleVenta.query.filter_by(
                    codigo_producto=codigo_anterior
                ).update({'codigo_producto': codigo}, synchronize_session=False)

            db.session.commit()
            flash(f'Producto "{codigo}" actualizado correctamente.', 'success')
            return redirect(_vista())

        # ---- ELIMINAR ----
        if action == 'eliminar':
            producto = Producto.query.get(id_prod) if id_prod else None
            if not producto:
                flash('El producto a eliminar no existe.', 'danger')
                return redirect(_vista())

            n_ent = Entrada.query.filter_by(codigo_producto=producto.codigo).count()
            n_det = DetalleVenta.query.filter_by(codigo_producto=producto.codigo).count()
            if n_ent or n_det:
                mensaje = (f'No se puede eliminar "{producto.codigo}": tiene '
                           f'{n_ent} compra(s) y {n_det} venta(s) registradas.')
                if n_det:
                    mensaje += ' Anula primero la venta desde Ventas.'
                if n_ent:
                    mensaje += ' Borra la compra desde Compras.'
                flash(mensaje, 'danger')
                return redirect(_vista())

            nombre_borrado = producto.codigo
            db.session.delete(producto)
            db.session.commit()
            flash(f'Producto "{nombre_borrado}" eliminado.', 'warning')
            return redirect(_vista())

        # ---- CREAR (por defecto) ----
        error = _validar_datos(upc, codigo, descripcion)
        if error:
            flash(error, 'danger')
            return redirect(_vista())

        nuevo = Producto(
            codigo=codigo,
            nombre=descripcion,
            codigo_barras=upc if upc else None,
        )
        db.session.add(nuevo)
        db.session.commit()

        flash(f'Producto "{codigo}" agregado correctamente.', 'success')
        return redirect(_vista())

    productos = _productos_unicos()
    return render_template('registros.html', productos=productos)


@registros_bp.route('/excel')
@login_required
def registros_excel():
    if current_user.rol != 'admin':
        flash('Solo el administrador puede exportar los registros.', 'danger')
        return redirect(url_for('main.index'))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    productos = _productos_unicos()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Registros Nuevos'

    encabezados = ['UPC', 'Codigo', 'Descripcion']
    fill = PatternFill(start_color='FE9F43', end_color='FE9F43', fill_type='solid')
    font_blanco = Font(bold=True, color='FFFFFF', size=11)
    borde = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )
    alineacion_c = Alignment(horizontal='center', vertical='center')

    for col, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.fill = fill
        celda.font = font_blanco
        celda.alignment = alineacion_c
        celda.border = borde

    fila = 2
    for p in productos:
        datos = [
            p.codigo_barras if p.codigo_barras else '',
            p.codigo,
            p.nombre,
        ]
        for col, valor in enumerate(datos, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.border = borde
            if col == 1:
                celda.alignment = alineacion_c
        fila += 1

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 70
    ws.auto_filter.ref = f'A1:C{fila - 1}'
    ws.freeze_panes = 'A2'

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    wb.close()

    return send_file(
        salida,
        as_attachment=True,
        download_name='registros_tech_house.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
