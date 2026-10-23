import openpyxl

EXCEL_PATH = r'G:\Mi unidad\SISTEMA TECH-HOUSE\TECH HOUSE.xlsm'
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

# --- INVENTARIO ---
ws = wb['Inventario']
print('=== INVENTARIO: Celdas con posibles formulas (None/0) ===')
inv_issues = 0
for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
    if row[0] and str(row[0]).startswith('PROD'):
        issues = []
        if row[2] is None: issues.append('costo=None')
        if row[2] == 0: issues.append('costo=0')
        if row[3] is None: issues.append('margen=None')
        if row[3] == 0: issues.append('margen=0')
        if row[4] is None: issues.append('precio=None')
        if row[4] == 0: issues.append('precio=0')
        if row[5] is None: issues.append('stock=None')
        if issues:
            inv_issues += 1
            print(f'  Row {i+3}: {row[0]} | {row[1]} | {", ".join(issues)}')
if inv_issues == 0:
    print('  OK - todas las celdas tienen valores')

# --- ENTRADA ---
ws = wb['Entrada']
print('\n=== ENTRADA: Celdas con posibles formulas ===')
ent_issues = 0
for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
    if row[0] and row[1]:
        issues = []
        if row[3] is None: issues.append('cantidad=None')
        if row[3] == 0: issues.append('cantidad=0')
        if row[4] is None: issues.append('costo=None')
        if row[4] == 0: issues.append('costo=0')
        if issues:
            ent_issues += 1
            print(f'  Row {i+3}: {row[1]} | {row[2]} | {", ".join(issues)}')
if ent_issues == 0:
    print('  OK - todas las celdas tienen valores')

# --- VENTAS ---
ws = wb['Ventas']
print('\n=== VENTAS: Celdas con posibles formulas ===')
ven_issues = 0
for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
    if row[0] and row[1]:
        issues = []
        if row[3] is None: issues.append('cantidad=None')
        if row[3] == 0: issues.append('cantidad=0')
        if row[4] is None: issues.append('total=None')
        if row[4] == 0: issues.append('total=0')
        if issues:
            ven_issues += 1
            print(f'  Row {i+3}: {row[1]} | {row[2]} | {", ".join(issues)}')
if ven_issues == 0:
    print('  OK - todas las celdas tienen valores')

# --- REPORTE ---
ws = wb['Reporte']
print('\n=== REPORTE: Celdas con posibles formulas ===')
rep_issues = 0
for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
    if row[0] and row[1]:
        issues = []
        if row[3] is None: issues.append('cantidad=None')
        if row[4] is None: issues.append('total=None')
        if row[5] is None: issues.append('ganancia=None')
        if issues:
            rep_issues += 1
            print(f'  Row {i+3}: {row[1]} | {row[2]} | {", ".join(issues)}')
if rep_issues == 0:
    print('  OK - todas las celdas tienen valores')

# --- Also read without data_only to detect formulas ---
print('\n=== Verificando formulas (sin data_only) ===')
wb2 = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=False)

ws_inv = wb2['Inventario']
formula_count = 0
for row in ws_inv.iter_rows(min_row=3, max_row=35):
    for cell in row:
        if cell.value and str(cell.value).startswith('='):
            formula_count += 1
            print(f'  Inventario {cell.coordinate}: {cell.value}')
print(f'  Total formulas en Inventario: {formula_count}')

ws_ent = wb2['Entrada']
formula_count = 0
for row in ws_ent.iter_rows(min_row=3, max_row=145):
    for cell in row:
        if cell.value and str(cell.value).startswith('='):
            formula_count += 1
            if formula_count <= 10:
                print(f'  Entrada {cell.coordinate}: {cell.value}')
if formula_count > 10:
    print(f'  ... y {formula_count - 10} mas')
print(f'  Total formulas en Entrada: {formula_count}')

ws_ven = wb2['Ventas']
formula_count = 0
for row in ws_ven.iter_rows(min_row=3, max_row=40):
    for cell in row:
        if cell.value and str(cell.value).startswith('='):
            formula_count += 1
            if formula_count <= 5:
                print(f'  Ventas {cell.coordinate}: {cell.value}')
print(f'  Total formulas en Ventas: {formula_count}')

ws_rep = wb2['Reporte']
formula_count = 0
for row in ws_rep.iter_rows(min_row=3, max_row=35):
    for cell in row:
        if cell.value and str(cell.value).startswith('='):
            formula_count += 1
            if formula_count <= 5:
                print(f'  Reporte {cell.coordinate}: {cell.value}')
print(f'  Total formulas en Reporte: {formula_count}')

wb.close()
wb2.close()
